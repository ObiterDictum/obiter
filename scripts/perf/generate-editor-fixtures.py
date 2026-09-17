"""Generate the deterministic DOCX fixtures editor-open measurement uses.

Real-toolchain provenance: every fixture is produced by python-docx, so it is
genuine Word-compatible OOXML rather than a hand-built zip. Fixtures are wholly
synthetic and fictional — no client matter data, no provider text.

The prose is generated from a fixed word pool with a seeded LCG, so the same
revision produces the same parts every time and DEFLATE sees a realistic ratio
(real English prose compresses about 2.5-4x; the upload guard rejects over 64x).
Zip entry timestamps record the save time, so the container bytes are not
themselves stable; every part inside is.

Every fixture carries a table, hard line breaks and bold/italic runs, and
citation-bearing paragraphs, so the model parser sees more than one story
shape. Sizes are chosen to sit well inside the 25 MiB upload cap.

Usage:
  python3 scripts/generate-editor-fixtures.py --out <dir>
"""
from __future__ import annotations

import argparse
import os
import re
import sys
import zipfile

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
except ImportError:  # pragma: no cover - operator instruction, not a test path
    sys.exit("python-docx is required: pip install python-docx")

WORDS = (
    "agreement lease underlease covenant boundary fence hedge drainage easement "
    "notice schedule valuation comparables surveyor report inspection instruction "
    "possession completion tranche retention drawdown assignment assignee trustee "
    "administrator liquidator guarantor surety indemnity warranty representation "
    "clause recital appendix schedule annex deed variation supplement novation "
    "landlord tenant lessee licensor licensee mortgagor mortgagee charge debenture "
    "consideration premium rent arrears forfeiture waiver estoppel rectification "
    "misrepresentation negligence nuisance trespass encroachment prescriptive user "
    "covenant covenant restrictive positive reservation exception proviso term "
    "condition precedent subsequent determination expert arbitrator adjudicator "
    "mediator tribunal court judge master registrar recorder district circuit "
    "appeal permission authority citation judgment order direction injunction "
    "declaration specific performance rectification rescission restitution damages "
    "interest costs indemnity contribution apportionment quantum mitigation "
    "causation remoteness foreseeability reliance inducement affirmation waiver "
    "land registry title plan filed transfer charge disposition registered "
    "unregistered possessory freehold leasehold tenancy licence profit covenant "
    "boundary exact line party structure wall fence ditch bank stream watercourse "
    "accessway right of way footpath bridleway easement quasi-easement prescription "
    "conveyance transfer deed instrument executed attested witness signature "
    "counterpart escrow completion undertaking covenant undertaking covenant "
).split()

FIRST = [
    "Cartwright", "Whitfield", "Ashcombe", "Draycott", "Halloway", "Pemberton",
    "Stannard", "Ravenscroft", "Thackeray", "Wickenden", "Marchetti", "Okonkwo",
]
TITLE = ["Mr", "Ms", "Mrs", "Dr", "Professor"]
HEADS = ["Background", "The lease", "The boundary dispute", "Quantum", "Costs",
         "Evidence", "Submissions", "Authorities"]


class Lcg:
    """Deterministic 31-bit LCG; a library seed would still be version-bound."""

    def __init__(self, seed: int) -> None:
        self.state = seed & 0x7FFFFFFF

    def next(self, modulus: int) -> int:
        self.state = (1103515245 * self.state + 12345) & 0x7FFFFFFF
        return (self.state >> 8) % modulus


def sentence(rng: Lcg, target: int) -> str:
    parts: list[str] = []
    length = 0
    while length < target:
        word = WORDS[rng.next(len(WORDS))]
        parts.append(word)
        length += len(word) + 1
    if rng.next(9) == 0:
        name = FIRST[rng.next(len(FIRST))]
        parts.insert(rng.next(len(parts) + 1),
                     f"{TITLE[rng.next(len(TITLE))]} {name}")
    text = " ".join(parts)
    return text[0].upper() + text[1:] + (". " if rng.next(4) else "? ")


def paragraph_text(rng: Lcg, words: int) -> str:
    out = ""
    while len(out.split()) < words:
        out += sentence(rng, 8 + rng.next(18))
    return out.strip()


def add_body(doc: Document, rng: Lcg, count: int) -> None:
    for i in range(count):
        p = doc.add_paragraph()
        words = 25 + rng.next(60)
        if i % 4 == 0:
            lead = p.add_run(paragraph_text(rng, 4) + " ")
            lead.bold = True
        run = p.add_run(paragraph_text(rng, words))
        if i % 6 == 1:
            run.italic = True
        if i % 9 == 7:
            p.add_run(" at [2019] EWCA Civ 1402, [2020] UKSC 14 and "
                      "[2018] EWHC 331 ")
            p.add_run().add_break(WD_BREAK.LINE)
            p.add_run(paragraph_text(rng, 20))


def add_table(doc: Document, rng: Lcg, rows: int) -> None:
    table = doc.add_table(rows=rows, cols=3)
    table.style = "Table Grid"
    for r in range(rows):
        for c in range(3):
            table.cell(r, c).text = paragraph_text(rng, 6)[:110]


def build(out_path: str, *, seed: int, sections: int, body_per_section: int,
          table_rows: int) -> dict:
    rng = Lcg(seed)
    doc = Document()
    for s in range(sections):
        heading = doc.add_heading(HEADS[s % len(HEADS)], level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_body(doc, rng, body_per_section)
        add_table(doc, rng, table_rows)
    doc.save(out_path)

    with zipfile.ZipFile(out_path) as z:
        uncompressed = sum(item.file_size for item in z.infolist())
        parts = len(z.infolist())
        xml = z.read("word/document.xml").decode("utf-8")
    return {
        "compressedBytes": os.path.getsize(out_path),
        "uncompressedBytes": uncompressed,
        "parts": parts,
        "paragraphs": len(re.findall(r"<w:p[ >]", xml)),
        "runs": len(re.findall(r"<w:r>", xml)),
        "hardBreaks": xml.count("<w:br/>"),
        "tables": xml.count("<w:tbl>"),
    }


SIZES = [
    ("small", dict(seed=101, sections=1, body_per_section=16, table_rows=4)),
    ("medium", dict(seed=202, sections=4, body_per_section=88, table_rows=14)),
    ("large", dict(seed=303, sections=6, body_per_section=300, table_rows=44)),
]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    os.makedirs(args.out, exist_ok=True)
    for name, kwargs in SIZES:
        path = os.path.join(args.out, f"{name}.docx")
        info = build(path, **kwargs)
        print(
            f"{name:7s} {info['compressedBytes']:>9,} B compressed "
            f"{info['uncompressedBytes']:>10,} B uncompressed | "
            f"{info['paragraphs']:>5} paragraphs {info['runs']:>6} runs "
            f"{info['hardBreaks']:>4} breaks {info['tables']:>3} tables "
            f"{info['parts']:>3} parts"
        )


if __name__ == "__main__":
    main()
